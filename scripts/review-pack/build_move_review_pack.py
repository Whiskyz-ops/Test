#!/usr/bin/env python3
"""Build the CA/CPA review pack for docs/CROSS_BORDER_MOVE_PLAN.md.

Reads the plan's law tables (residency, treaty, US/India rules, DTAA
application, planning, liabilities, gap sweeps, accounts, decision guides,
state exit) and its open law questions, adds the proposed case list, and
writes an Excel workbook reviewers fill in. Engineering content (engine
nodes, form fields, phasing) is left out on purpose.

    python3 scripts/review-pack/build_move_review_pack.py

Output: docs/review/CROSS_BORDER_MOVE_REVIEW_PACK.xlsx. Rebuild it whenever
the plan changes; reviewer answers live in the copy sent out, not here.
"""
import re
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
PLAN = ROOT / "docs" / "CROSS_BORDER_MOVE_PLAN.md"
OUT = ROOT / "docs" / "review" / "CROSS_BORDER_MOVE_REVIEW_PACK.xlsx"

FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SUB = Font(name=FONT, size=11, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
FILL_HEAD = PatternFill("solid", fgColor="1F3A5F")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_EXAMPLE = PatternFill("solid", fgColor="FFF2CC")
FILL_BAND = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

VERDICTS = ["Agree", "Agree with change", "Disagree", "Not sure / needs research", "Not my area"]


# ---------------------------------------------------------------- parsing
def clean(text: str) -> str:
    """Strip Markdown and plan-internal cross-references."""
    t = text.replace("\\|", "|")
    t = re.sub(r"`([^`]*)`", r"\1", t)
    t = t.replace("**", "")
    t = re.sub(r"(?<![\w*])\*([^*\n]+)\*(?![\w*])", r"\1", t)
    t = re.sub(r"~~([^~]*)~~", r"\1", t)
    t = re.sub(r"§12 item (\d+)", r"open question Q-\1", t)
    # plan section refs are §0–§13 with an optional .n; IRC refs (§121,
    # §72(t), §25B, §1014) are longer or followed by "(" / a letter.
    plan_ref = r"§(?:1[0-3]|[0-9])(?:\.[0-9]+)?(?![0-9A-Za-z(])"
    t = re.sub(r"\s*\((?:see |as in )?" + plan_ref + r"(?:[,;–\- ]+" + plan_ref + r")*\)", "", t)
    t = re.sub(plan_ref, lambda m: "plan " + m.group(0), t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def plan_tables(lines):
    """Yield (heading, header_cells, rows) for every Markdown table."""
    heading, i = "", 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("## ") or line.startswith("### "):
            heading = line.lstrip("#").strip()
        if line.startswith("|") and (i == 0 or not lines[i - 1].startswith("|")):
            j = i
            while j < len(lines) and lines[j].startswith("|"):
                j += 1
            block = lines[i:j]
            split = lambda l: [c.strip() for c in re.split(r"(?<!\\)\|", l.strip())[1:-1]]
            yield heading, split(block[0]), [split(l) for l in block[2:]]
            i = j
            continue
        i += 1


def open_questions(lines):
    start = next(i for i, l in enumerate(lines) if l.startswith("## 12. Law items"))
    end = next(i for i, l in enumerate(lines) if i > start and l.startswith("## "))
    items, cur = [], None
    for l in lines[start + 1:end]:
        m = re.match(r"^(\d+)\. (.*)", l)
        if m:
            if cur:
                items.append(cur)
            cur = [int(m.group(1)), m.group(2)]
        elif cur and l.startswith("   "):
            cur[1] += " " + l.strip()
    if cur:
        items.append(cur)
    return [(f"Q-{n}", QUESTION_TEXT.get(f"Q-{n}", clean(t))) for n, t in items]


def area_for(ref: str, default: str) -> str:
    if "-US-" in ref:
        return "US (CPA)"
    if "-IN-" in ref:
        return "India (CA)"
    return default


def build_rules(lines):
    rules = []  # (ref, scenario, area, topic, rule, confidence)
    # §3.1 / §3.2 — written as prose in the plan; restated here one rule per row.
    rules += [
        ("R-1", "Both", "India (CA)", "No split year in India", "India's residential status (ROR / RNOR / NR) applies to the whole financial year (1 April – 31 March). A mid-year move does not split the Indian year.", "Certain"),
        ("R-2", "Moving to the US", "US (CPA)", "US residency start — substantial presence test", "Residency starts on the first day present in the US in the year the test is met, ignoring up to 10 days of earlier presence where the person had a closer connection to India.", "Certain"),
        ("R-3", "Moving to the US", "US (CPA)", "US residency start — green card", "Residency starts on the first day present in the US as a lawful permanent resident.", "Certain"),
        ("R-4", "Moving to the US", "US (CPA)", "First-year choice §7701(b)(4)", "Available when the substantial presence test isn't met in the arrival year but is met the next year: needs 31 consecutive days of presence and presence on 75% of the days from then to year-end. The return waits until next year's test is met (Form 4868 extension).", "Certain"),
        ("R-5", "Returning to India", "US (CPA)", "US residency end date", "Residency ends on the last day present if, for the rest of the year, the person has a tax home in India and a closer connection to India, and is not a US resident the next year. A statement is attached to the return; without that showing, the person is resident to 31 December. (Reg. 301.7701(b)-4)", "Likely (statement requirement to confirm)"),
        ("R-6", "Returning to India", "US (CPA)", "Green card end of residency", "US residency continues until the green card is surrendered (Form I-407) or abandonment is determined.", "Certain"),
        ("R-7", "Both", "US (CPA)", "US citizens", "US citizens never stop being US tax residents; a move changes FEIE and foreign-tax-credit planning, not the scope of US tax.", "Certain"),
        ("R-8", "Both", "India (CA)", "India residency tests (s.6)", "Resident if in India 182+ days in the FY, or 60+ days in the FY and 365+ days over the previous 4 FYs. The 60-day limb becomes 182 days for an Indian citizen leaving India for employment abroad (Explanation 1(a)) and for a citizen or person of Indian origin on a visit; for a visitor with Indian income over ₹15L it becomes 120 days.", "Certain"),
        ("R-9", "Returning to India", "India (CA)", "RNOR", "RNOR if non-resident in 9 of the previous 10 FYs, or in India 729 days or fewer over the previous 7 FYs. Also RNOR: deemed residents under s.6(1A) and visitors resident under the 120-day rule.", "Certain"),
        ("R-10", "Returning to India", "India (CA)", "Returning for good is not a visit", "A person returning to India permanently is not 'on a visit', so the plain 60-day limb (with 365 days over 4 years) applies.", "Likely (practitioner consensus)"),
        ("R-11", "Both", "US (CPA)", "Dual-status income attribution", "In a dual-status year each item is taxed under the rules of the period in which it was received or accrued (resident period: worldwide; non-resident period: US-source / effectively connected only), not by day-count.", "Likely (Reg. 1.871-13)"),
    ]

    a_dt = b_dt = t_n = ra_n = ia_n = dg_n = 0
    for heading, header, rows in plan_tables(lines):
        h = heading
        if h.startswith("3.3"):
            for r in rows:
                t_n += 1
                rules.append((f"T-{t_n}", "Both", "Treaty (CA + CPA)", clean(r[0]), clean(r[1]), clean(r[2])))
        elif h.startswith(("4.3", "4.4", "5.3", "5.4")):
            scen = "Returning to India" if h.startswith("4") else "Moving to the US"
            for r in rows:
                rules.append((r[0], scen, area_for(r[0], "Both"), clean(r[1]), clean(r[2]), clean(r[3])))
        elif h.startswith(("4.5", "5.6")):
            scen = "Returning to India" if h.startswith("4") else "Moving to the US"
            for r in rows:
                if h.startswith("4"):
                    a_dt += 1; ref = f"A-DT-{a_dt}"
                else:
                    b_dt += 1; ref = f"B-DT-{b_dt}"
                rules.append((ref, scen, "Treaty (CA + CPA)", clean(r[0]), f"{clean(r[2])} (DTAA Art. {clean(r[1])})", "Likely"))
        elif h.startswith("5.5"):
            for r in rows:
                rules.append((r[0], "Moving to the US", "Both", clean(r[1]), clean(r[2]), clean(r[3])))
        elif h.startswith("8.2"):
            for r in rows:
                rules.append((r[0], "Both", "Both", clean(r[1]), clean(r[2]), clean(r[3])))
        elif h.startswith(("8.3", "8.9")):
            for r in rows:
                if r[0] in ("G-27", "G-28", "G-29", "G-30"):  # product / covered elsewhere / out of scope
                    continue
                rules.append((r[0], "Both", "Both", clean(r[1]), clean(r[2]), clean(r[3])))
        elif h.startswith("8.5") and header[0] == "Account" and len(header) == 6:
            for r in rows:
                ra_n += 1
                text = (f"Contributions after moving: {clean(r[1])} | US withdrawals: {clean(r[2])} | "
                        f"India: {clean(r[3])} | What to do: {clean(r[4])}")
                rules.append((f"RA-{ra_n}", "Both", "Both", f"US account: {clean(r[0])}", text, clean(r[5])))
        elif h.startswith("8.5") and header[0] == "Account":
            for r in rows:
                ia_n += 1
                text = f"On becoming NR in India: {clean(r[1])} | US once resident: {clean(r[2])} | What to do: {clean(r[3])}"
                rules.append((f"IA-{ia_n}", "Moving to the US", "Both", f"Indian account: {clean(r[0])}", text, clean(r[4])))
        elif h.startswith("8.7"):
            for r in rows:
                dg_n += 1
                rules.append((f"DG-{dg_n}", "Both", "Both", f"{clean(r[0])} — options: {clean(r[1])}", clean(r[2]), clean(r[3])))

    rules += [
        ("ST-1", "Returning to India", "US (CPA)", "State domicile-break evidence", "Surrender the driver's licence; cancel car registration or sell the car; cancel voter registration; remove any homestead exemption; give up or rent out any home in the state; change address with employer, banks, brokers and the IRS (Form 8822); move or close safe-deposit boxes; end memberships relevant to domicile tests.", "Likely"),
        ("ST-2", "Returning to India", "US (CPA)", "California", "Closest-connection test and the employment-abroad safe harbour. RSUs and options stay CA-taxable by CA workdays. 3⅓% withholding on a non-resident's CA property sale.", "Likely"),
        ("ST-3", "Returning to India", "US (CPA)", "New York", "Statutory resident at 183 days plus a permanent place of abode. Residency audits of former residents are common; keep day logs.", "Certain"),
        ("ST-4", "Returning to India", "US (CPA)", "State returns after leaving", "Final part-year return for the move year, then non-resident returns for state-source income (rent from the old home, its sale, RSUs, deferred comp not protected by 4 U.S.C. §114). Washington taxes capital gains despite no income tax.", "Likely"),
        ("ST-5", "Returning to India", "US (CPA)", "Retirement income protected (4 U.S.C. §114)", "A former state can't tax most retirement income received after the person leaves (qualified plans, IRAs, government plans, and deferred comp paid over 10+ years).", "Certain"),
        ("ST-6", "Moving to the US", "US (CPA)", "State entry", "Part-year resident return for the arrival year. Most states give no credit for Indian tax.", "Likely"),
    ]
    # Plan rows that mention engine internals, restated for a reviewer.
    out = []
    for ref, scen, area, topic, rule, conf in rules:
        if ref in RULE_TEXT:
            rule = RULE_TEXT[ref]
        rule = rule.replace(" (built for full-year 1040-NR)", "")
        out.append((ref, scen, area, topic, rule, conf))
    return out


RULE_TEXT = {
    "T-7": "Pensions and annuities are taxable only in the country of residence. Our current reading, already reviewed by a CPA: for an India resident filing Form 1040-NR, periodic payments (including SEPP payments while the schedule holds) are exempt under Art. 20(1); lump sums are taxed as FDAP at 30% plus §72(t); a broken SEPP is treated as a lump sum.",
    "T-8": "Art. 21(2): standard deduction on Form 1040-NR. Art. 21(1): payments from outside the US for maintenance and study are exempt; Art. 1(4)(b) keeps this for a resident alien but not for a citizen or green-card holder.",
    "T-9": "Teaching and research pay is exempt for two years from the date of arrival when the person was resident in India immediately before. Art. 1(4)(b) keeps it for a resident alien, not for a citizen or green-card holder.",
    "A-US-10": "For a full-year Form 1040-NR filer resident in India: periodic payments exempt (Art. 20(1)); lump sums taxed as FDAP at 30% plus §72(t); a broken SEPP treated as a lump sum. Proposed for the move year: withdrawals paid before the residency end date fall in the resident period (Form 1040, §72(t)); those paid after follow the 1040-NR treaty rule.",
    "A-DT-2": "The US wins the tie-breaker for the overlap window (permanent home), or India credits US tax on US-earned pay. Both routes should be shown. (DTAA Art. 4, 16, 25)",
    "L-4": "US cancellation-of-debt income arises when debt is settled or forgiven around the move. How it is sourced for a non-resident needs legal confirmation.",
    "G-9": "ESPP: India taxes the discount at purchase; the US taxes at sale (qualifying or disqualifying disposition). Workday sourcing applies when the offering period straddles the move.",
    "G-14": "A returning resident's PAN must be linked to Aadhaar or it becomes inoperative (higher TDS, refunds blocked). NRIs are exempt.",
    "G-20": "Once ROR, India taxes gains on US shares. Computing them needs each lot's purchase date and cost, converted to INR at the purchase-date rate.",
    "G-23": "Exchange rates: India uses Rule 115 (SBI TT buying rate); the US uses yearly average or spot rates. The move year makes this harder because the two periods use different rates.",
}

QUESTION_TEXT = {
    "Q-4": "DTAA Art. 20 vs Art. 23 for lump-sum withdrawals, and §72(t) under the treaty: settled for full-year 1040-NR filers by an earlier CPA review. Still open: India's side of a lump sum (Art. 23(3) lets the source country tax, so India credits the US tax).",
    "Q-11": "Whether a person returning to India for good is 'on a visit' for the 182 / 120-day extension of the 60-day limb. Our reading: not on a visit, so the plain 60-day + 365-day test applies.",
}


# ---------------------------------------------------------------- cases
CASES = [
    # id, direction, US status, example move, India status in the move FY, household, first/repeat, key facts, why separate, persona
    ("C-A1", "Returning to India", "Visa (H-1B)", "15 Aug 2026", "RNOR", "Single", "First", "9 years in US; RSUs vesting after departure; 401(k) and Roth; keeps US rental condo; NRE/NRO accounts", "Core returnee case", "A1"),
    ("C-A2", "Returning to India", "Green card, long-term resident (10 yrs)", "15 Aug 2026", "RNOR", "Single", "First", "Files I-407; net worth near §877A thresholds", "§877A covered-expatriate tests; Form 8854", "A2"),
    ("C-A3", "Returning to India", "Green card, under 8 years", "Jun 2026", "RNOR", "Single", "First", "Files I-407", "No §877A; residency ends at surrender", ""),
    ("C-A4", "Returning to India", "Green card kept (re-entry permit)", "Jun 2026", "RNOR", "Single", "First", "Lives in India, keeps card", "Stays US resident on worldwide income (G-3)", ""),
    ("C-A5", "Returning to India", "US citizen (OCI)", "Aug 2026", "RNOR", "Single", "First", "Relocating permanently; Indian employer", "US scope unchanged; FEIE vs FTC; EPF international worker", "A3"),
    ("C-A6", "Returning to India", "Visa (H-1B)", "Aug 2026", "ROR (no RNOR)", "Single", "First", "Only 3 years abroad", "Overlap Apr–Aug taxed by both; Art. 4/25", "A4"),
    ("C-A7", "Returning to India", "Visa (H-1B)", "Feb 2026", "NR in FY 2025-26, RNOR from FY 2026-27", "Single", "First", "Return in Jan–Mar", "Move year spans two Indian years", ""),
    ("C-A8", "Returning to India", "Visa (H-1B)", "Nov 2026", "NR for FY 2026-27 (unless 60+365), RNOR next FY", "Single", "First", "Late-year return; few prior visits", "India status for the move FY may be NR", ""),
    ("C-A9", "Returning to India", "Visa (H-1B)", "Aug 2026", "RNOR", "Single", "First", "Keeps US ties; no closer-connection claim", "US resident all year (residency end date not supported)", ""),
    ("C-A10", "Returning to India", "Both spouses on visas", "Aug 2026", "RNOR (both)", "Couple, same date", "First", "Filed MFJ in prior years", "No MFJ in dual-status year without election", ""),
    ("C-A11", "Returning to India", "Both spouses on visas", "Spouse 1 Jun, spouse 2 Dec 2026", "RNOR / NR", "Couple, different dates", "First", "One spouse stays to finish work", "Different residency periods per spouse", ""),
    ("C-A12", "Returning to India", "US-citizen spouse + visa spouse", "Aug 2026", "RNOR (visa spouse)", "Couple, same date", "First", "Citizen spouse stays a US taxpayer", "§6013(g) option for the non-citizen spouse", ""),
    ("C-A13", "Returning to India", "Visa (H-1B)", "Aug 2026", "RNOR", "Family with US-born children", "First", "Children are US citizens with Indian accounts", "Children's own US filings, FBAR, PFIC (G-4)", ""),
    ("C-A14", "Returning to India", "F-1 student", "Jun 2026", "RNOR", "Single", "First", "Exempt individual throughout", "US non-resident all along; no dual-status year", ""),
    ("C-A15", "Returning to India", "Visa (H-1B)", "Aug 2026", "ROR", "Single", "Repeat", "Returned to India once before within 4 years", "No RNOR on a repeat return (G-31)", ""),
    ("C-A16", "Returning to India", "Visa / green card", "Aug 2026", "RNOR", "Single", "First", "Consultant keeping US clients from India", "Business income sourcing; no US SE tax (G-11)", ""),
    ("C-A17", "Returning to India", "Green card", "Aug 2026", "RNOR", "Single", "First", "Owns an S-corp", "S election ends on becoming non-resident alien", ""),
    ("C-A18", "Returning to India", "Green card / citizen, age 62+", "Aug 2026", "RNOR", "Couple, same date", "First", "Retiring: pension annuity, IRA, Social Security, RMDs soon", "Retirement-income treaty treatment; Medicare", ""),
    ("C-B1", "Moving to the US", "Visa (H-1B)", "15 Aug 2026", "NR (left for employment)", "Single", "First", "Indian mutual funds; NRO FD", "SPT not met: full-year NRA vs first-year choice", "B1"),
    ("C-B2", "Moving to the US", "Visa (H-1B) + H-4 spouse", "Feb 2026 (spouse Apr)", "ROR in FY 2025-26, NR after", "Couple, different dates", "First", "Spouse has Indian rental income", "Dual-status vs §6013(h); Jan–Mar move", "B2"),
    ("C-B3", "Moving to the US", "F-1 student", "Aug 2026", "NR", "Single", "First", "Funding from family in India", "Exempt individual; Art. 21; Form 8843", "B3"),
    ("C-B4", "Moving to the US", "Green card immigrant", "Nov 2026", "ROR (182+ days in India)", "Couple, same date", "First", "Flat, EPF/PPF, ULIP", "India ROR in departure FY; overlap Nov–Mar", "B4"),
    ("C-B5", "Moving to the US", "Visa (H-1B)", "May 2026", "NR", "Single", "First", "SPT met in arrival year", "Dual-status from arrival date (most common arrival)", ""),
    ("C-B6", "Moving to the US", "Green card immigrant", "Jun 2026", "NR", "Single", "First", "Immigrant visa", "Residency from first day as permanent resident", ""),
    ("C-B7", "Moving to the US", "J-1 researcher / teacher", "Aug 2026", "NR", "Single", "First", "University appointment", "Art. 22 two-year exemption", ""),
    ("C-B8", "Moving to the US", "L-1 transferee", "Jul 2026", "NR", "Single", "First", "Paid partly by Indian employer (secondment)", "Indian payroll for a US resident; Art. 16", ""),
    ("C-B9", "Moving to the US", "Visa (H-1B)", "Jul 2026", "NR", "Single", "First", "Owns 20% of an Indian company and an LLP share", "Forms 5471/8865, CFC rules from arrival", ""),
    ("C-B10", "Moving to the US", "Visa (H-1B)", "Jul 2026", "NR", "Single", "First", "Member of an HUF", "Expected output: 'not covered' warning (G-5)", ""),
    ("C-B11", "Moving to the US", "Both spouses on visas", "Jul 2026", "NR (both)", "Couple, same date", "First", "Both employed in India before", "First-year choice + §6013(h) joint election", ""),
    ("C-B12", "Moving to the US", "Visa (H-1B)", "Jul 2026", "NR", "Couple, spouse stays in India", "First", "Spouse has Indian salary", "MFS vs §6013(g) for the non-resident spouse", ""),
    ("C-B13", "Moving to the US", "Visa (H-1B)", "Mar 2026", "NR", "Single", "Repeat", "Lived in the US 2 years ago", "SPT counts prior-year days; earlier PFIC history", ""),
    ("C-B14", "Moving to the US", "US citizen raised in India", "Jul 2026", "NR", "Single", "First", "Has always filed US returns from India", "US scope unchanged; India departure rules", ""),
    ("C-B15", "Moving to the US", "Visa / green card", "Year after move", "RNOR risk in a later year", "Single", "First", "Visits India 130 days; Indian income over ₹15L", "Becoming resident in India again (G-2)", ""),
    ("C-B16", "Moving to the US", "Visa (H-1B)", "Jul 2026", "NR", "Single", "First", "Keeps working remotely for Indian employer for 2 months", "Indian pay received while US resident (G-38)", ""),
]


# ---------------------------------------------------------------- workbook
def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEAD, FILL_HEAD, Alignment(wrap_text=True, vertical="center"), BORDER


def body(cell, bold=False, fill=None):
    cell.font = F_BOLD if bold else F_BODY
    cell.alignment = WRAP
    cell.border = BORDER
    if fill:
        cell.fill = fill


def widths(ws, ws_widths):
    for col, w in ws_widths.items():
        ws.column_dimensions[col].width = w


def add_verdict_dv(ws, col_letter, first_row, last_row):
    dv = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True,
                        error="Choose a verdict from the list.", errorTitle="Verdict")
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{first_row}:{col_letter}{last_row}")


def main():
    lines = PLAN.read_text().split("\n")
    rules = build_rules(lines)
    questions = open_questions(lines)
    commit = subprocess.run(["git", "log", "-1", "--format=%h", "--", str(PLAN)], cwd=ROOT,
                            capture_output=True, text=True).stdout.strip() or "unknown"

    wb = Workbook()

    # ---------------- Rules
    ws = wb.active
    ws.title = "Rules"
    hdr = ["Ref", "Scenario", "Reviewer", "Topic", "Rule as drafted", "Our confidence", "Priority",
           "Verdict", "Correct rule (if different)", "Authority / citation", "Comments", "Reviewed by", "Date"]
    ws.append(hdr)
    style_header(ws, 1, len(hdr))
    for i, (ref, scen, area, topic, rule, conf) in enumerate(rules, start=2):
        ws.append([ref, scen, area, topic, rule, conf, None, None, None, None, None, None, None])
        ws.cell(row=i, column=7).value = (f'=IF(ISNUMBER(SEARCH("Verify",F{i})),"1 High",'
                                          f'IF(ISNUMBER(SEARCH("Likely",F{i})),"2 Medium","3 Low"))')
        for c in range(1, 14):
            body(ws.cell(row=i, column=c), bold=(c == 1), fill=FILL_INPUT if c >= 8 else None)
        ws.cell(row=i, column=13).number_format = "yyyy-mm-dd"
    last_rule = len(rules) + 1
    add_verdict_dv(ws, "H", 2, last_rule)
    widths(ws, {"A": 9, "B": 15, "C": 15, "D": 26, "E": 70, "F": 16, "G": 10, "H": 16, "I": 40, "J": 26, "K": 30, "L": 14, "M": 11})
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:M{last_rule}"

    # ---------------- Open questions
    wq = wb.create_sheet("Open questions")
    hdr_q = ["Ref", "Question", "Reviewer", "Answer", "Authority / citation", "Comments", "Answered by", "Date"]
    wq.append(hdr_q)
    style_header(wq, 1, len(hdr_q))
    india_q = {1, 12, 13, 14, 15, 18, 20, 21, 23, 25}
    us_q = {2, 3, 7, 9, 17, 19, 22}
    for i, (ref, q) in enumerate(questions, start=2):
        n = int(ref.split("-")[1])
        who = "India (CA)" if n in india_q else "US (CPA)" if n in us_q else "Both"
        wq.append([ref, q, who, None, None, None, None, None])
        for c in range(1, 9):
            body(wq.cell(row=i, column=c), bold=(c == 1), fill=FILL_INPUT if c >= 4 else None)
        wq.cell(row=i, column=8).number_format = "yyyy-mm-dd"
    last_q = len(questions) + 1
    widths(wq, {"A": 7, "B": 70, "C": 14, "D": 50, "E": 26, "F": 30, "G": 14, "H": 11})
    wq.freeze_panes = "B2"

    # ---------------- Cases
    wc = wb.create_sheet("Cases")
    hdr_c = ["Case", "Direction", "US status", "Example move date", "India status (move FY)", "Household",
             "First or repeat move", "Key facts", "Why it is a separate case", "Plan persona",
             "Keep as a separate case?", "Expected outcome notes / missing variants", "Reviewed by"]
    wc.append(hdr_c)
    style_header(wc, 1, len(hdr_c))
    for i, row in enumerate(CASES, start=2):
        wc.append(list(row) + [None, None, None])
        for c in range(1, 14):
            body(wc.cell(row=i, column=c), bold=(c == 1), fill=FILL_INPUT if c >= 11 else None)
    last_c = len(CASES) + 1
    dv = DataValidation(type="list", formula1='"Yes,No - merge,No - drop"', allow_blank=True)
    wc.add_data_validation(dv)
    dv.add(f"K2:K{last_c}")
    note_row = last_c + 2
    wc.cell(row=note_row, column=1, value="Missing cases: add them in the rows below, starting with ID C-NEW1.").font = F_NOTE
    for r in range(note_row + 1, note_row + 6):
        wc.cell(row=r, column=1, value=f"C-NEW{r - note_row}")
        for c in range(1, 14):
            body(wc.cell(row=r, column=c), bold=(c == 1), fill=FILL_INPUT if c >= 2 else None)
    widths(wc, {"A": 8, "B": 16, "C": 22, "D": 16, "E": 22, "F": 18, "G": 11, "H": 36, "I": 36, "J": 8, "K": 14, "L": 40, "M": 14})
    wc.freeze_panes = "B2"

    # ---------------- Guide sample
    wg = wb.create_sheet("Guide sample")
    wg.append(["Part", "Draft text (example finding: NRE account on return to India)", "Right level of detail?", "Correction / comments"])
    style_header(wg, 1, 4)
    guide = [
        ("What changes", "The client becomes resident in India under FEMA on 15 Aug 2026. The Axis NRE account (₹19L) has to be converted."),
        ("Why", "FEMA. The NRE interest exemption (s.10(4)(ii)) applies only to a person resident outside India under FEMA. [Likely; ITA 2025 section to map]"),
        ("What to do", "1. Client, on return: tell the bank in writing; redesignate the NRE savings account as a resident account, or move the foreign-currency part to an RFC account.\n2. Client: NRE deposits may run to maturity; interest from the return date is taxable. [Likely; check the bank's terms]\n3. Specialist: include interest from the return date in Indian income. While RNOR, check FCNR/RFC money, whose interest stays exempt.\n4. Specialist: if the client was a US resident for part of the year, report that period's interest on the US return."),
        ("Forms and documents", "The bank's redesignation form; passport copy with the arrival stamp; the bank's interest certificate split at the return date."),
        ("Deadline", "15 Aug 2026 (the return date)."),
        ("Cost of not acting", "Indian tax plus interest on interest wrongly treated as exempt. The FEMA penalty needs legal confirmation."),
        ("Doesn't apply if", "The client is back only on a visit and keeps FEMA non-resident status — specialist judgement."),
        ("Confidence", "Likely"),
    ]
    for i, (part, text) in enumerate(guide, start=2):
        wg.append([part, text, None, None])
        for c in range(1, 5):
            body(wg.cell(row=i, column=c), bold=(c == 1), fill=FILL_INPUT if c >= 3 else None)
    dvg = DataValidation(type="list", formula1='"Right,Too little,Too much,Wrong"', allow_blank=True)
    wg.add_data_validation(dvg)
    dvg.add(f"C2:C{len(guide) + 1}")
    widths(wg, {"A": 20, "B": 90, "C": 16, "D": 50})

    # ---------------- Context
    wx = wb.create_sheet("Context")
    r = 1
    wx.cell(row=r, column=1, value="Decision windows used in the rules").font = F_SUB
    r += 1
    win_hdr = ["Direction", "Window", "US status", "India status", "Typical effect"]
    for c, v in enumerate(win_hdr, 1):
        wx.cell(row=r, column=c, value=v)
    style_header(wx, r, 5)
    windows = next(rows for h, hd, rows in plan_tables(lines) if h.startswith("8.4"))
    for row in windows:
        r += 1
        for c, v in enumerate(row, 1):
            body(wx.cell(row=r, column=c, value=clean(v)))
    r += 2
    wx.cell(row=r, column=1, value="Example timelines").font = F_SUB
    r += 1
    for c, v in enumerate(["Example", "US tax year 2026", "India FY 2025-26 (to 31 Mar 2026)", "India FY 2026-27", "Months resident in both"], 1):
        wx.cell(row=r, column=c, value=v)
    style_header(wx, r, 5)
    for row in [
        ("Returns to India 15 Aug 2026 (H-1B, 9 yrs in US)", "Resident 1 Jan–15 Aug; non-resident after", "NR", "RNOR for the whole year", "Apr–Aug 2026"),
        ("Arrives in US 10 Feb 2026 (H-1B)", "Non-resident to 9 Feb; resident from 10 Feb", "ROR for the whole year", "NR for the whole year", "10 Feb–31 Mar 2026"),
    ]:
        r += 1
        for c, v in enumerate(row, 1):
            body(wx.cell(row=r, column=c, value=v))
    r += 2
    wx.cell(row=r, column=1, value="Confidence tags").font = F_SUB
    for tag, meaning in [("Certain", "Settled law, routinely applied."), ("Likely", "Strong reading, widely applied, but with edge cases or fact-dependence."), ("Verify", "Must be checked against the statute, treaty or ITA 2025 renumbering before we build it.")]:
        r += 1
        body(wx.cell(row=r, column=1, value=tag), bold=True)
        body(wx.cell(row=r, column=2, value=meaning))
    widths(wx, {"A": 40, "B": 40, "C": 30, "D": 28, "E": 50})

    # ---------------- Read me (first sheet)
    wr = wb.create_sheet("Read me", 0)
    widths(wr, {"A": 30, "B": 22, "C": 22, "D": 22, "E": 40})
    r = 1
    wr.cell(row=r, column=1, value="India ⇄ US cross-border moves — professional review pack").font = F_TITLE
    r += 1
    wr.cell(row=r, column=1, value=f"Built from docs/CROSS_BORDER_MOVE_PLAN.md (commit {commit}). Law as understood on 2 Oct 2026 for US tax year 2026 and India FY 2026-27. Drafts for review, not advice.").font = F_NOTE
    r += 2
    paras = [
        ("What this is", "WISING is building support for clients who move between India and the US: returning NRIs and people moving to the US. Before anything is built, every tax rule we plan to apply needs an Indian CA and a US CPA to confirm it, correct it, or say we're missing something."),
        ("What we ask", "1. Rules: give a verdict on each rule; if it's wrong, write the correct rule and the authority.\n2. Open questions: answer the questions we couldn't settle ourselves.\n3. Cases: tell us whether each case really behaves differently, and add any case we've missed.\n4. Guide sample: tell us whether the step-by-step guidance is at the right level of detail."),
        ("Who reviews what", "The 'Reviewer' column says who each row is for: India (CA), US (CPA), Treaty (both), or Both. Please answer only what's in your area; choose 'Not my area' otherwise."),
        ("How to fill it in", "Only the yellow cells. Verdict, Keep-as-a-case and Right-level cells have drop-down lists. Sort by 'Priority' to see the rules tagged Verify first. References like 'plan §8.4' point to the full plan document, sent separately; you don't need it to answer."),
        ("Section numbers", "Indian sections use ITA 1961 numbering (s.6, s.10(4)(ii), s.115H) unless a row says otherwise. Mapping them to the Income-tax Act 2025 is open question Q-1."),
    ]
    for label, text in paras:
        body(wr.cell(row=r, column=1, value=label), bold=True)
        wr.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        c = wr.cell(row=r, column=2, value=text)
        body(c)
        wr.row_dimensions[r].height = 15 * (1 + text.count("\n") + len(text) // 95)
        r += 1
    r += 1
    wr.cell(row=r, column=1, value="Example of a completed row (format only — not a real answer)").font = F_SUB
    r += 1
    for c, v in enumerate(["Ref", "Verdict", "Correct rule (if different)", "Authority / citation", "Comments"], 1):
        wr.cell(row=r, column=c, value=v)
    style_header(wr, r, 5)
    r += 1
    for c, v in enumerate(["G-13", "Agree with change", "Also note: 1040-NR with wages subject to withholding is due 15 April.", "Instructions for Form 1040-NR (2026), 'When to file'", "Rest is correct."], 1):
        body(wr.cell(row=r, column=c, value=v), fill=FILL_EXAMPLE)
    r += 2
    wr.cell(row=r, column=1, value="Progress").font = F_SUB
    r += 1
    progress = [
        ("Rules in this pack", f"=COUNTA(Rules!A2:A{last_rule})"),
        ("Rules with a verdict", f"=COUNTA(Rules!H2:H{last_rule})"),
        ("  Agree", f'=COUNTIF(Rules!H2:H{last_rule},"Agree")'),
        ("  Agree with change", f'=COUNTIF(Rules!H2:H{last_rule},"Agree with change")'),
        ("  Disagree", f'=COUNTIF(Rules!H2:H{last_rule},"Disagree")'),
        ("  Not sure / needs research", f'=COUNTIF(Rules!H2:H{last_rule},"Not sure / needs research")'),
        ("High-priority (Verify) rules", f'=COUNTIF(Rules!G2:G{last_rule},"1 High")'),
        ("Open questions", f"=COUNTA('Open questions'!A2:A{last_q})"),
        ("Open questions answered", f"=COUNTA('Open questions'!D2:D{last_q})"),
        ("Cases", f"=COUNTA(Cases!A2:A{last_c})"),
        ("Cases reviewed", f"=COUNTA(Cases!K2:K{last_c})"),
    ]
    for label, formula in progress:
        body(wr.cell(row=r, column=1, value=label))
        body(wr.cell(row=r, column=2, value=formula))
        r += 1

    # No cached values are written (openpyxl can't compute them), so make
    # Excel / Google Sheets / Numbers calculate every formula on open.
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{OUT.relative_to(ROOT)}: {len(rules)} rules, {len(questions)} questions, {len(CASES)} cases")


if __name__ == "__main__":
    main()
